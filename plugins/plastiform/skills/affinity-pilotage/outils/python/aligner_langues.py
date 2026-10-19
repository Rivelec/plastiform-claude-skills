#!/usr/bin/env python3
"""aligner_langues.py — aligne les textes d'une langue cible sur ceux de la langue de référence (FR) et signale les
problèmes de traduction détectables automatiquement, classés par gravité.

Usage : python3 aligner_langues.py REFERENCE.json CIBLE.json [--lexique LEXIQUE.xlsx|.csv|.json] [--md rapport.md]
        [--json rapport.json] [--qcm questions.json] [--ratio 1.35] [--langue DE] [--col-source FR] [--col-cible 3]
  REFERENCE.json, CIBLE.json : produits par affinity/inventaire_textes.js (un par fichier de langue).
  --lexique : termes validés (xlsx : 1re feuille ; csv ; json [{source, cible}]). Colonnes : celle du français et celle
              de --langue (en-tête « (DE) » ou « Allemand », de préférence « Validé »), ou --col-source / --col-cible
              (numéro ou texte de l'en-tête). « A / B » = variantes admises ; « (s) », « (n) », « (ex. …) » ignorés.
  --qcm     : écrit un squelette de questions (qcm.py generer) pour les écarts au lexique et les textes non traduits.
Appariement : même planche et même chemin d'objet si possible, sinon le bloc le plus proche (centre) sur la même
planche. Contrôles : non traduit (texte identique), nombres différents (valeurs, unités), bloc manquant ou en plus,
terme du lexique non respecté, texte nettement plus long (risque de débordement).
"""
import argparse, csv, json, os, re, sys, unicodedata

def creer_dossier(p):
    """Crée le dossier d'un fichier de sortie s'il n'existe pas (ex. « revue/rapport.md »), renvoie le chemin."""
    d = os.path.dirname(os.path.abspath(p))
    os.makedirs(d, exist_ok=True)
    return p


GRAV = ("CRITIQUE", "MAJEUR", "MINEUR")


def norm(t):
    return re.sub(r"\s+", " ", (t or "").replace("\u2029", " ").replace("\u2028", " ")).strip()


def sans_accents(t):
    """minuscules, sans accents, traits d'union = espaces (« Standard-Mischdüse » = « standard mischduse »)."""
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"\s+", " ", t.replace("-", " "))


def codes(t):
    """Références et codes (AC-011, CA-F20-2, DM-50…) : repères stables d'une langue à l'autre."""
    return set(re.findall(r"\b[A-Z]{1,4}-[A-Z0-9][A-Z0-9-]*\b", t))


def nombres(t):
    """Nombres d'un texte, séparateurs décimaux unifiés (1,24 = 1.24 ; 1 234 = 1234) ; les entiers de 0 à 12, souvent
    écrits en toutes lettres dans une autre langue (« 2 » / « zwei »), sont ignorés."""
    t = re.sub(r"(?<![A-Za-z\d])(\d{1,3})[\u00a0\u202f ](?=\d{3}\b)", r"\1", t)   # 1 234 → 1234, mais pas « M25 900 ml »
    return sorted(x for x in (y.replace(",", ".") for y in re.findall(r"\d+(?:[.,]\d+)?", t)) if "." in x or int(x) > 12)


def centre(b):
    return (b[0] + b[2] / 2, b[1] + b[3] / 2) if b else (0, 0)


NOMS_LANGUES = {"DE": "allemand", "EN": "anglais", "ES": "espagnol", "IT": "italien", "PT": "portugais", "NL": "néerlandais", "PL": "polonais", "FR": "français"}


ARTICLES = r"^(?:der|die|das|den|dem|des|ein|eine|einen|the|a|an|el|la|los|las|un|una|le|les|il|lo|gli)\s+"


def nettoyer_terme(t):
    """« Mallette(s) » → « Mallette » ; « Lot (ex: lot de 2) » → « Lot » ; « Dose(n) » → « Dose » ;
    « Der Maßabdruck » → « Maßabdruck » (article retiré)."""
    t = re.sub(r"\s*\([^)]*\)", "", str(t)).strip(" .;:")
    return re.sub(ARTICLES, "", t, flags=re.I).strip()


def charger_proscrits(spec):
    """Termes à proscrire dans la langue cible : « fichier.xlsx:2 » = 2e feuille ; colonnes détectées par l'en-tête
    (« proscrire », « interdit », « éviter » / « valid », « standard ») sur l'une des 5 premières lignes."""
    if not spec:
        return []
    import openpyxl
    chemin, _, feuille = spec.partition(":")
    ws = openpyxl.load_workbook(chemin, read_only=True, data_only=True).worksheets[int(feuille or 1) - 1]
    lignes = [list(l) for l in ws.iter_rows(values_only=True)]
    for k, l in enumerate(lignes[:5]):
        ent = [str(x or "").lower() for x in l]
        ip = next((i for i, x in enumerate(ent) if "proscri" in x or "interdit" in x or "éviter" in x), None)
        iv = next((i for i, x in enumerate(ent) if "valid" in x or "standard" in x), None)
        if ip is not None and iv is not None:
            out = []
            for l2 in lignes[k + 1:]:
                if len(l2) > max(ip, iv) and l2[ip] and l2[iv]:
                    for mot in re.split(r"\s*/\s*|\s*;\s*", str(l2[ip])):
                        mot = nettoyer_terme(mot)
                        if len(mot) >= 4:
                            out.append((mot, str(l2[iv]).strip()))
            return out
    return []


def colonne(entetes, choix, langue, defaut):
    if choix:
        if str(choix).isdigit():
            return int(choix) - 1
        return next(i for i, x in enumerate(entetes) if str(choix).lower() in x)
    nom = NOMS_LANGUES.get((langue or "").upper(), "")
    cands = [i for i, x in enumerate(entetes) if (langue and "(%s)" % langue.lower() in x) or (nom and nom in x)]
    valides = [i for i in cands if "valid" in entetes[i]]
    return (valides or cands or [defaut])[0]


def charger_lexique(p, langue_src="FR", langue_cible=None, col_src=None, col_cible=None):
    if not p:
        return []
    paires = []
    if p.lower().endswith(".json"):
        j = json.load(open(p, encoding="utf-8"))
        for e in (j if isinstance(j, list) else j.get("termes", [])):
            paires.append((e.get("source") or e.get("fr"), e.get("cible")))
    else:
        if p.lower().endswith(".csv"):
            lignes = list(csv.reader(open(p, encoding="utf-8-sig")))
        else:
            import openpyxl
            ws = openpyxl.load_workbook(p, read_only=True, data_only=True).worksheets[0]
            lignes = [list(l) for l in ws.iter_rows(values_only=True)]
        ent = [str(x or "").lower() for x in lignes[0]]
        isrc = colonne(ent, col_src, langue_src, 0)
        icib = colonne(ent, col_cible, langue_cible, 1)
        for l in lignes[1:]:
            if l and len(l) > max(isrc, icib) and l[isrc] and l[icib]:
                paires.append((str(l[isrc]), str(l[icib])))
    out = []
    for s, c in paires:
        if not s or not c:
            continue
        sources = [nettoyer_terme(x) for x in re.split(r"\s*/\s*|\s*;\s*", str(s)) if nettoyer_terme(x)]
        cands = [nettoyer_terme(x) for x in re.split(r"\s*/\s*|\s*;\s*", str(c)) if nettoyer_terme(x)]   # « A / B » : variantes admises
        for src in sources:
            if len(src) >= 3 and cands:
                out.append((src, cands))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("reference"); ap.add_argument("cible")
    ap.add_argument("--lexique"); ap.add_argument("--md"); ap.add_argument("--json"); ap.add_argument("--qcm")
    ap.add_argument("--ratio", type=float, default=1.35)
    ap.add_argument("--langue", help="langue cible (DE, EN, ES…) pour choisir la colonne du lexique")
    ap.add_argument("--col-source"); ap.add_argument("--col-cible")
    ap.add_argument("--proscrits", help="XLSX[:n° de feuille] : colonne « à proscrire » (termes interdits dans la cible) et colonne « validé / standard »")
    a = ap.parse_args()
    def charger(p):
        L = [t for t in json.load(open(p, encoding="utf-8")) if norm(t.get("texte"))]
        for t in L:   # tableaux : cellules séparées (sinon « 96 » et « 192 » se collent en « 96192 »)
            if t.get("cellules"):
                t["texte"] = " | ".join(norm(c) for c in t["cellules"])
        return L
    R, T = charger(a.reference), charger(a.cible)
    lex = sorted(charger_lexique(a.lexique, "FR", a.langue, a.col_source, a.col_cible), key=lambda x: -len(x[0]))
    proscrits = charger_proscrits(a.proscrits)
    ecarts_lex, ecarts_pro = {}, {}
    par_chemin = {(t["si"], t["chemin"]): t for t in T}
    libres = set(range(len(T)))
    index = {id(t): i for i, t in enumerate(T)}
    paires, pb = [], []
    for r in R:
        t = par_chemin.get((r["si"], r["chemin"]))
        if t is not None and (abs(centre(t.get("boite"))[0] - centre(r.get("boite"))[0]) > 300 or index[id(t)] not in libres):
            t = None   # même chemin mais objet différent : on cherche par position
        if t is None:   # bloc le plus proche sur la même planche, départagé par les codes communs (AC-011…)
            cr = centre(r.get("boite")); cod = codes(r["texte"]); best, smin, dbest, cbest = None, 1e18, 0, 0
            for i in libres:
                u = T[i]
                if u["si"] != r["si"]:
                    continue
                cu = centre(u.get("boite")); d = ((cu[0] - cr[0]) ** 2 + (cu[1] - cr[1]) ** 2) ** 0.5
                commun = len(cod & codes(u["texte"]))
                score = d - 400 * commun
                if score < smin:
                    best, smin, dbest, cbest = u, score, d, commun
            if best is not None and (dbest < 250 or cbest > 0) and not (cod and codes(best["texte"]) and not cbest):
                t = best
        if t is None:
            pb.append(("MAJEUR", r["page"], "bloc sans équivalent dans la cible", norm(r["texte"]), "", bool(r.get("cellules"))))
            continue
        libres.discard(index[id(t)])
        paires.append((r, t))
    for i in sorted(libres):
        u = T[i]
        pb.append(("MINEUR", u["page"], "bloc en plus dans la cible (sans équivalent en référence)", "", norm(u["texte"]), bool(u.get("cellules"))))
    for r, t in paires:
        s, c = norm(r["texte"]), norm(t["texte"])
        tab = bool(t.get("cellules"))   # tableau : texte = cellules jointes par « | » (pas de règle « exact » possible)
        lettres = re.sub(r"[^A-Za-zÀ-ÿ]", "", s)
        code = re.fullmatch(r"[A-Z0-9][A-Z0-9 ./_+-]*", s) is not None
        if s == c and len(lettres) > 3 and not code and t.get("visible", True):
            pb.append(("MAJEUR", t["page"], "non traduit (texte identique à la référence)", s, c, tab))
        nr, nc = nombres(s), nombres(c)
        if nr != nc and (nr or nc):
            manq = sorted(set(nr) - set(nc)); plus = sorted(set(nc) - set(nr))
            if manq or plus:
                pb.append(("MAJEUR", t["page"], "nombres différents (manquants : %s ; en plus : %s)" % (", ".join(manq) or "—", ", ".join(plus) or "—"), s, c, tab))
        if len(s) > 20 and len(c) > a.ratio * len(s):
            pb.append(("MINEUR", t["page"], "texte %.0f %% plus long que la référence (risque de débordement)" % (100.0 * len(c) / len(s) - 100), s, c, tab))
        ss, cc = sans_accents(s), sans_accents(c)
        couverts = []   # zones déjà couvertes par un terme plus long (« prise d'empreinte » avant « empreinte »)
        for src, cands in lex:
            m = re.search(r"(?<![\w])" + re.escape(sans_accents(src)) + r"(?:s|x)?(?![\w])", ss)
            if not m or any(a0 <= m.start() and m.end() <= b0 for a0, b0 in couverts):
                continue
            couverts.append((m.start(), m.end()))
            if not any(sans_accents(x) in cc for x in cands):
                ecarts_lex.setdefault((src, tuple(cands)), []).append((t["page"], s, c, tab))
        for mot, alt in proscrits:
            if re.search(r"(?<![\w])" + re.escape(sans_accents(mot)), cc):
                ecarts_pro.setdefault((mot, alt), []).append((t["page"], s, c, tab))
    # termes proscrits regroupés ; une règle dont le terme proscrit figure dans le terme validé (« Mischdüsen » →
    # « Spezial-Mischdüsen ») dépend du contexte : MINEUR, à vérifier
    for (mot, alt), occ in sorted(ecarts_pro.items(), key=lambda x: -len(x[1])):
        contexte = sans_accents(mot) in sans_accents(alt)
        pb.append(("MINEUR" if contexte else "MAJEUR", occ[0][0], "terme proscrit « %s » (validé : « %s »%s) dans %d bloc(s) (pages %s)" % (
            mot, alt, ", règle de contexte : vérifier" if contexte else "", len(occ), ", ".join(str(p) for p in sorted({o[0] for o in occ})[:12])), occ[0][1], occ[0][2], occ[0][3]))
    # écarts au lexique regroupés par terme : MAJEUR s'ils touchent un titre ou un libellé court, MINEUR sinon
    for (src, cands), occ in sorted(ecarts_lex.items(), key=lambda x: -len(x[1])):
        courts = [o for o in occ if len(o[1]) <= 60]
        ex = (courts or occ)[0]
        pb.append(("MAJEUR" if courts else "MINEUR", ex[0], "lexique : « %s » → « %s » absent de %d bloc(s)%s (pages %s)" % (
            src, " / ".join(cands), len(occ), " dont %d titre(s) ou libellé(s)" % len(courts) if courts else "",
            ", ".join(str(p) for p in sorted({o[0] for o in occ})[:12])), ex[1], ex[2], ex[3]))
    ordre = {g: i for i, g in enumerate(GRAV)}
    pb.sort(key=lambda x: (ordre[x[0]], x[1] if isinstance(x[1], int) else 999))
    cpt = {g: sum(1 for x in pb if x[0] == g) for g in GRAV}
    nom_r, nom_c = os.path.basename(a.reference), os.path.basename(a.cible)
    md = ["# Alignement %s → %s" % (nom_r, nom_c), "",
          "%d bloc(s) apparié(s) ; %d CRITIQUE, %d MAJEUR, %d MINEUR%s" % (len(paires), cpt["CRITIQUE"], cpt["MAJEUR"], cpt["MINEUR"], " ; lexique : %d terme(s)" % len(lex) if lex else ""), ""]
    for g in GRAV:
        items = [x for x in pb if x[0] == g]
        if items:
            md += ["## " + g, "", "| page | constat | référence | cible |", "|---|---|---|---|"]
            court = lambda t: (t[:120] + "…" if len(t) > 120 else t).replace("|", "/")
            md += ["| %s | %s | %s | %s |" % (x[1], x[2], court(x[3]), court(x[4])) for x in items] + [""]
    texte = "\n".join(md) + "\n"
    if a.md:
        open(creer_dossier(a.md), "w", encoding="utf-8").write(texte)
    if a.json:
        json.dump({"compte": cpt, "paires": len(paires), "problemes": [dict(zip(("gravite", "page", "constat", "reference", "cible"), x)) for x in pb]},
                  open(creer_dossier(a.json), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    if a.qcm:
        qs = []
        for k, x in enumerate([x for x in pb if x[2].startswith("lexique") or x[2].startswith("non traduit")][:150]):
            q = {"id": "Q-%03d" % (k + 1), "gravite": x[0], "theme": "Lexique" if x[2].startswith("lexique") else "Non traduit",
                 "titre": x[2], "ou": "p. %s" % x[1], "actuel": x[4], "options": [], "contextes": [{"etiquette": "Réf.", "page": x[1], "texte": x[3]}, {"etiquette": "Cible", "page": x[1], "texte": x[4], "cible": True}]}
            if x[5]:   # tableau : « exact » ne peut pas viser une cellule → décision à traduire en règle à la main
                q["quoi"] = "Tableau : indiquer la cellule à corriger et sa nouvelle valeur (règle écrite ensuite à la main)."
            else:
                q["regle"] = {"f": x[4], "exact": 1}   # texte COMPLET du bloc (« exact » compare le bloc entier)
            qs.append(q)
        json.dump({"config": {"cle": "alignement-" + nom_c, "titre": "Traduction : points à trancher", "intro": "Proposer la bonne traduction (option « Autre ») ou garder le texte actuel."}, "questions": qs},
                  open(creer_dossier(a.qcm), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(texte if len(texte) < 7000 else texte[:7000] + "\n… (complet : %s)" % (a.md or a.json or "--md"))
    return 1 if cpt["CRITIQUE"] else 0


if __name__ == "__main__":
    sys.exit(main())
