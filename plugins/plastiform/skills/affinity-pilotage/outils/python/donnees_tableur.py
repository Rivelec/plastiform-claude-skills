#!/usr/bin/env python3
"""donnees_tableur.py — passe le fichier de données centralisées (tarifs.json, donnees.json…) en tableur Excel pour
le modifier confortablement, puis reporte les valeurs modifiées dans le JSON (avant appliquer_donnees.js).

  python3 donnees_tableur.py exporter DONNEES.json TABLEUR.xlsx
      une ligne par valeur : n° de bloc, page, contexte, ligne, valeur actuelle, NOUVELLE VALEUR (à remplir) ;
      colonnes de repérage (planche, x, y) en gris, à ne pas modifier.
  python3 donnees_tableur.py importer TABLEUR.xlsx DONNEES.json [--appliquer]
      sans --appliquer : liste des changements (simulation) ; avec : écrit DONNEES.json (copie de l'ancien en
      DONNEES.avant_<date>.json). Une cellule « nouvelle valeur » vide = valeur inchangée.
"""
import argparse, datetime, json, os, shutil, sys

def creer_dossier(p):
    """Crée le dossier d'un fichier de sortie s'il n'existe pas (ex. « revue/rapport.md »), renvoie le chemin."""
    d = os.path.dirname(os.path.abspath(p))
    os.makedirs(d, exist_ok=True)
    return p



def valeurs(t):
    if "valeurs" in t: return t["valeurs"], "valeurs"
    if "prix" in t: return t["prix"], "prix"
    return [l.get("prix") for l in t.get("lignes", [])], "lignes"


def exporter(a):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    d = json.load(open(a.json, encoding="utf-8"))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Données"
    ent = ["bloc", "page", "contexte", "ligne", "référence", "valeur actuelle", "NOUVELLE VALEUR", "planche", "x", "y"]
    ws.append(ent)
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
    gris = PatternFill("solid", fgColor="EEEEEE"); jaune = PatternFill("solid", fgColor="FFF2CC")
    for i, t in enumerate(d.get("tableaux", [])):
        vals, cle = valeurs(t)
        for k, v in enumerate(vals):
            ref = (t["lignes"][k].get("ref", "") + " " + t["lignes"][k].get("designation", "")).strip() if cle == "lignes" else ""
            ws.append([i, t.get("page"), t.get("contexte", ""), k + 1, ref, v, None, t.get("planche"), t.get("x"), t.get("y")])
            r = ws.max_row
            ws.cell(r, 7).fill = jaune
            for col in (8, 9, 10):
                ws.cell(r, col).fill = gris
    for col, w in zip("ABCDEFGHIJ", (6, 6, 28, 6, 40, 16, 18, 8, 8, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    note = wb.create_sheet("Mode d'emploi")
    for l in ["Remplir uniquement la colonne « NOUVELLE VALEUR », au format affiché (ex. 1 234,00 €).",
              "Une cellule vide = valeur inchangée. Ne pas trier ni supprimer de ligne ; ne pas modifier planche, x, y.",
              "Ensuite : python3 donnees_tableur.py importer <ce fichier> <données.json> (simulation), puis --appliquer,",
              "puis appliquer_donnees.js dans Affinity (simulation puis application), puis l'export."]:
        note.append([l])
    note.column_dimensions["A"].width = 110
    wb.save(creer_dossier(a.xlsx))
    print("%s : %d ligne(s) de valeurs" % (a.xlsx, ws.max_row - 1))


def importer(a):
    import openpyxl
    d = json.load(open(a.json, encoding="utf-8"))
    ws = openpyxl.load_workbook(a.xlsx, data_only=True)["Données"]
    blocs = d.get("tableaux", [])
    chg, err = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[6] in (None, ""):
            continue
        i, k, nv = int(row[0]), int(row[3]) - 1, str(row[6]).strip()
        if i >= len(blocs):
            err.append("bloc %s inexistant" % i); continue
        t = blocs[i]
        if (t.get("planche"), t.get("x"), t.get("y")) != (row[7], row[8], row[9]):
            err.append("bloc %s : repérage modifié dans le tableur (lignes triées ?) — ignoré" % i); continue
        vals, cle = valeurs(t)
        if k >= len(vals):
            err.append("bloc %s ligne %s inexistante" % (i, k + 1)); continue
        if str(vals[k]) == nv:
            continue
        chg.append((t.get("page"), t.get("contexte", ""), k + 1, vals[k], nv))
        if cle == "lignes":
            t["lignes"][k]["prix"] = nv
        else:
            t[cle][k] = nv
    for p, ctx, k, av, nv in chg:
        print("p.%s %s ligne %s : %s → %s" % (p, ctx, k, av, nv))
    for e in err:
        print("ERREUR " + e)
    if not a.appliquer:
        print("SIMULATION : %d changement(s). Relancer avec --appliquer pour écrire %s." % (len(chg), a.json)); return
    if chg:
        sauv = os.path.splitext(a.json)[0] + ".avant_%s.json" % datetime.datetime.now().strftime("%Y%m%d-%H%M")
        shutil.copyfile(a.json, sauv)
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(d, f, ensure_ascii=False, indent=1)
        print("%d changement(s) écrit(s) dans %s (ancien : %s)" % (len(chg), a.json, os.path.basename(sauv)))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    e = sp.add_parser("exporter"); e.add_argument("json"); e.add_argument("xlsx"); e.set_defaults(f=exporter)
    i = sp.add_parser("importer"); i.add_argument("xlsx"); i.add_argument("json"); i.add_argument("--appliquer", action="store_true"); i.set_defaults(f=importer)
    a = ap.parse_args(); a.f(a)


if __name__ == "__main__":
    main()
